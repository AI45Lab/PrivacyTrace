"""Population, city, education, religion, migration, and household targets."""

from __future__ import annotations

import csv
import gzip
import io
import json
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl

from .common import (
    download,
    normalize,
    overlap_fraction,
    parse_age_range,
    read_zipped_csv,
)
from .constants import (
    AGE_GROUPS,
    BASE_ETHNICITY,
    COUNTRY_CODES,
    EDUCATION_INDICATORS,
    EDUCATION_LEVELS,
    ETHNICITIES,
    IMS_ORIGIN_TO_ETHNICITY,
    RELATIONSHIP_MAP,
    SEX_MAP,
    UNDATA_URL,
    WORLD_BANK_API,
    WUP_CITY_URL,
)


def load_wpp(path: Path) -> dict[str, Any]:
    age_sex: dict[tuple[str, int, str, str], float] = defaultdict(float)
    age5_sex: dict[tuple[str, int, int, int, str], float] = defaultdict(float)
    total: dict[tuple[str, int], float] = defaultdict(float)
    total_sex: dict[tuple[str, int, str], float] = defaultdict(float)
    adult: dict[tuple[str, int], float] = defaultdict(float)
    supported = {info["iso3"]: country for country, info in COUNTRY_CODES.items()}
    with gzip.open(path, "rt", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            if row["ISO3_code"] not in supported or int(row["Time"]) not in {
                2020,
                2025,
            }:
                continue
            country = supported[row["ISO3_code"]]
            year = int(row["Time"])
            low = int(row["AgeGrpStart"])
            span = int(row["AgeGrpSpan"])
            high = low + span - 1
            male = float(row["PopMale"]) * 1000
            female = float(row["PopFemale"]) * 1000
            age5_sex[(country, year, low, high, "male")] += male
            age5_sex[(country, year, low, high, "female")] += female
            total[(country, year)] += male + female
            total_sex[(country, year, "male")] += male
            total_sex[(country, year, "female")] += female
            adult_fraction = overlap_fraction(low, high, 18, 80)
            adult[(country, year)] += (male + female) * adult_fraction
            if year == 2025:
                for age_group, (target_low, target_high) in AGE_GROUPS.items():
                    fraction = overlap_fraction(low, high, target_low, target_high)
                    if fraction:
                        age_sex[(country, year, age_group, "male")] += male * fraction
                        age_sex[(country, year, age_group, "female")] += (
                            female * fraction
                        )
    return {
        "age_sex": age_sex,
        "age5_sex": age5_sex,
        "total": total,
        "total_sex": total_sex,
        "adult": adult,
    }


def build_population_targets(
    wpp: dict[str, Any],
) -> tuple[dict[str, float], dict[str, float]]:
    country_population = normalize(
        {country: wpp["adult"][(country, 2025)] for country in COUNTRY_CODES}
    )
    age_sex_population = {
        f"{age_group} | {sex}": sum(
            wpp["age_sex"][(country, 2025, age_group, sex)] for country in COUNTRY_CODES
        )
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    }
    return country_population, normalize(age_sex_population)


def build_country_age_sex_targets(wpp: dict[str, Any]) -> dict[str, float]:
    population = {
        f"{country} | {age_group} | {sex}": wpp["age_sex"][
            (country, 2025, age_group, sex)
        ]
        for country in COUNTRY_CODES
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    }
    return normalize(population)


def city_size_class(population: float) -> str:
    if population < 100_000:
        return "50k-100k"
    if population < 250_000:
        return "100k-250k"
    if population < 500_000:
        return "250k-500k"
    if population < 1_000_000:
        return "500k-1m"
    if population < 5_000_000:
        return "1m-5m"
    if population < 10_000_000:
        return "5m-10m"
    return "10m+"


def load_wup_city_catalog(
    path: Path,
    country_mass: dict[str, float],
) -> tuple[dict[str, Any], dict[str, float]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Data"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    index = {str(value): position for position, value in enumerate(header)}
    iso_to_country = {info["iso3"]: country for country, info in COUNTRY_CODES.items()}
    catalog: dict[str, list[dict[str, Any]]] = {
        country: [] for country in COUNTRY_CODES
    }
    for row in rows:
        iso3 = row[index["ISO3_Code"]]
        country = iso_to_country.get(str(iso3))
        population_thousands = row[index["2025"]]
        city_name = row[index["City_Name"]]
        if (
            country is None
            or city_name is None
            or population_thousands is None
            or float(population_thousands) < 50.0
        ):
            continue
        population = round(float(population_thousands) * 1000)
        catalog[country].append(
            {
                "name": str(city_name).strip(),
                "population_2025": population,
                "size_class": city_size_class(population),
            }
        )

    targets: dict[str, float] = {}
    for country, cities in catalog.items():
        if not cities:
            raise RuntimeError(f"WUP returned no 50,000+ cities for {country}.")
        cities.sort(key=lambda item: (-item["population_2025"], item["name"]))
        by_size: dict[str, float] = defaultdict(float)
        for city in cities:
            by_size[city["size_class"]] += city["population_2025"]
        conditional = normalize(by_size)
        for size_class, share in conditional.items():
            targets[f"{country} | {size_class}"] = country_mass[country] * share

    output = {
        "source": WUP_CITY_URL,
        "year": 2025,
        "minimum_population": 50_000,
        "countries": catalog,
    }
    return output, targets


def fetch_wdi_education(cache: Path) -> dict[str, Any]:
    cache_file = cache / "world_bank_education_attainment.json"
    if cache_file.exists():
        cached = json.loads(cache_file.read_text(encoding="utf-8"))
    else:
        cached = {}
    iso_to_country = {info["iso3"]: country for country, info in COUNTRY_CODES.items()}
    country_path = ";".join(iso_to_country)
    for label, indicator in EDUCATION_INDICATORS.items():
        if all(label in cached.get(country, {}) for country in COUNTRY_CODES):
            continue
        url = WORLD_BANK_API.format(
            countries=country_path,
            indicator=indicator,
        )
        with urllib.request.urlopen(url, timeout=240) as response:
            payload = json.load(response)
        rows = payload[1] if isinstance(payload, list) and len(payload) > 1 else []
        seen: set[str] = set()
        for row in rows:
            iso3 = str(row.get("countryiso3code", ""))
            country = iso_to_country.get(iso3)
            if country is None or country in seen or row.get("value") is None:
                continue
            cached.setdefault(country, {})[label] = {
                "year": int(row["date"]),
                "value": float(row["value"]) / 100.0,
                "indicator": indicator,
            }
            seen.add(country)
    cache_file.write_text(
        json.dumps(cached, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return cached


def build_education_targets(
    country_mass: dict[str, float],
    observations: dict[str, Any],
) -> tuple[dict[str, float], dict[str, Any]]:
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for country, mass in country_mass.items():
        values = {
            label: min(1.0, max(0.0, float(item["value"])))
            for label, item in observations[country].items()
        }
        primary = values["primary"]
        lower = min(primary, values["lower_secondary"])
        upper = min(lower, values["upper_secondary"])
        post = min(upper, values.get("post_secondary", 0.0))
        bachelor = min(upper - post, values["bachelor"])
        master = min(bachelor, values["master"])
        doctorate = min(master, values.get("doctorate", 0.0))
        shares = {
            "none": 1.0 - primary,
            "primary": primary - lower,
            "lower secondary": lower - upper,
            "upper secondary": upper - post - bachelor,
            "vocational": post,
            "bachelor": bachelor - master,
            "master": master - doctorate,
            "doctorate": doctorate,
        }
        shares = normalize(shares)
        for level in EDUCATION_LEVELS:
            targets[f"{country} | {level}"] = mass * shares[level]
        audit[country] = {
            "indicator_observations": observations[country],
            "education_level_shares": shares,
        }
    return targets, audit


def load_pew_religion(
    path: Path,
    country_mass: dict[str, float],
) -> tuple[dict[str, float], dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        name = next(
            item for item in archive.namelist() if item.endswith("(percentages).csv")
        )
        text = io.TextIOWrapper(archive.open(name), encoding="utf-8-sig")
        rows = list(csv.DictReader(text))
    m49_to_country = {info["m49"]: country for country, info in COUNTRY_CODES.items()}
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for row in rows:
        if row.get("Year") != "2020":
            continue
        try:
            country = m49_to_country[int(row["Countrycode"])]
        except (KeyError, TypeError, ValueError):
            continue
        shares = normalize(
            {
                "Christian": float(row["Christians"]),
                "Muslim": float(row["Muslims"]),
                "Unaffiliated": float(row["Religiously_unaffiliated"]),
                "Buddhist": float(row["Buddhists"]),
                "Hindu": float(row["Hindus"]),
                "Jewish": float(row["Jews"]),
                "Folk": float(row["Other_religions"]),
            }
        )
        for religion, share in shares.items():
            targets[f"{country} | {religion}"] = country_mass[country] * share
        audit[country] = {
            "source_country": row["Country"],
            "year": 2020,
            "religious_belief_shares": shares,
        }
    missing = sorted(set(COUNTRY_CODES) - set(audit))
    if missing:
        raise RuntimeError(f"Pew religion data missing countries: {missing}")
    return targets, audit


def load_ims(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Table 1"]
    code_to_country = {info["m49"]: country for country, info in COUNTRY_CODES.items()}
    destinations: dict[str, dict[str, Any]] = {}
    wanted_origins = set(code_to_country) | set(IMS_ORIGIN_TO_ETHNICITY) | {900}
    for row in sheet.iter_rows(min_row=12, values_only=True):
        destination_code = row[3]
        origin_code = row[6]
        if destination_code not in code_to_country or origin_code not in wanted_origins:
            continue
        country = code_to_country[destination_code]
        entry = destinations.setdefault(
            country,
            {
                "data_type": str(row[4] or ""),
                "origins": {},
                "origins_by_sex": {"male": {}, "female": {}},
            },
        )
        entry["origins"][int(origin_code)] = float(row[13] or 0)
        entry["origins_by_sex"]["male"][int(origin_code)] = float(row[20] or 0)
        entry["origins_by_sex"]["female"][int(origin_code)] = float(row[27] or 0)
    return destinations


def build_migration_targets(
    country_mass: dict[str, float],
    wpp: dict[str, Any],
    ims: dict[str, Any],
) -> tuple[dict[str, float], dict[str, Any]]:
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for destination, mass in country_mass.items():
        origins = ims[destination]["origins"]
        population = wpp["total"][(destination, 2020)]
        migrant_total = origins[900]
        native = max(0.0, population - migrant_total)
        supported = {
            origin: origins.get(info["m49"], 0.0)
            for origin, info in COUNTRY_CODES.items()
            if origin != destination
        }
        denominator = native + sum(supported.values())
        row = {origin: value / denominator for origin, value in supported.items()}
        row[destination] = native / denominator
        for origin in COUNTRY_CODES:
            targets[f"{destination} | {origin}"] = mass * row.get(origin, 0.0)
        basis = ims[destination]["data_type"]
        audit[destination] = {
            "ims_data_type": basis,
            "direct_for_birth_country": "B" in basis,
            "direct_for_citizenship_country": "C" in basis,
            "wpp_2020_population": round(population),
            "ims_2020_migrant_total": round(migrant_total),
            "supported_foreign_origin_total": round(sum(supported.values())),
            "conditioning": (
                f"native plus foreign origins in the supported "
                f"{len(COUNTRY_CODES)}-country universe"
            ),
        }
    return targets, audit


def build_migration_status_by_sex_targets(
    wpp: dict[str, Any],
    ims: dict[str, Any],
) -> tuple[dict[str, float], dict[str, Any]]:
    """Build current-country × sex × local/foreign status from IMS Table 1."""

    total_adult_population = sum(
        wpp["age_sex"][(country, 2025, age_group, sex)]
        for country in COUNTRY_CODES
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    )
    targets: dict[str, float] = {}
    audit: dict[str, Any] = {}
    for country in COUNTRY_CODES:
        audit[country] = {}
        for sex in ("female", "male"):
            origins = ims[country]["origins_by_sex"][sex]
            migrant_total = origins[900]
            native = max(
                0.0,
                wpp["total_sex"][(country, 2020, sex)] - migrant_total,
            )
            supported_foreign = sum(
                origins.get(info["m49"], 0.0)
                for origin, info in COUNTRY_CODES.items()
                if origin != country
            )
            conditional = normalize({"local": native, "foreign": supported_foreign})
            sex_mass = (
                sum(
                    wpp["age_sex"][(country, 2025, age_group, sex)]
                    for age_group in AGE_GROUPS
                )
                / total_adult_population
            )
            for status, share in conditional.items():
                targets[f"{country} | {sex} | {status}"] = sex_mass * share
            audit[country][sex] = {
                "wpp_2020_same_sex_population": round(
                    wpp["total_sex"][(country, 2020, sex)]
                ),
                "ims_2020_same_sex_migrant_total": round(migrant_total),
                "supported_foreign_origin_total": round(supported_foreign),
                "conditional_status_shares": conditional,
            }
    return targets, audit


def load_marital_rows(cache: Path) -> tuple[dict[str, Any], dict[str, int]]:
    country_rates: dict[str, Any] = {}
    years: dict[str, int] = {}
    for country, info in COUNTRY_CODES.items():
        url = UNDATA_URL.format(table=23, country_code=info["m49"])
        try:
            rows = read_zipped_csv(
                download(url, cache / f"undata_23_{info['m49']}.zip")
            )
        except Exception:
            continue
        rows = [
            row
            for row in rows
            if row.get("Year", "").isdigit()
            and row.get("Area") == "Total"
            and row.get("Sex") in SEX_MAP
        ]
        if not rows:
            continue
        latest = max(int(row["Year"]) for row in rows)
        years[country] = latest
        rows = [row for row in rows if int(row["Year"]) == latest]
        rates: dict[tuple[str, str], dict[str, float]] = {}
        for age_group, (target_low, target_high) in AGE_GROUPS.items():
            for source_sex, sex in SEX_MAP.items():
                counts = {status: 0.0 for status in set(RELATIONSHIP_MAP.values())}
                for row in rows:
                    if (
                        row["Sex"] != source_sex
                        or row["Marital status"] not in RELATIONSHIP_MAP
                    ):
                        continue
                    age_range = parse_age_range(row["Age"])
                    if not age_range:
                        continue
                    fraction = overlap_fraction(*age_range, target_low, target_high)
                    if fraction:
                        counts[RELATIONSHIP_MAP[row["Marital status"]]] += (
                            float(row["Value"]) * fraction
                        )
                rates[(age_group, sex)] = normalize(counts)
        country_rates[country] = rates
    return country_rates, years


def build_relationship_targets(
    age_sex_mass: dict[str, float],
    wpp: dict[str, Any],
    country_rates: dict[str, Any],
) -> dict[str, float]:
    targets: dict[str, float] = {}
    statuses = ["single", "in relationship", "married", "divorced", "widowed"]
    for age_group in AGE_GROUPS:
        for sex in ("female", "male"):
            weighted = {status: 0.0 for status in statuses}
            denominator = 0.0
            for country in COUNTRY_CODES:
                if country not in country_rates:
                    continue
                population = wpp["age_sex"][(country, 2025, age_group, sex)]
                denominator += population
                for status in statuses:
                    weighted[status] += population * country_rates[country][
                        (age_group, sex)
                    ].get(status, 0.0)
            conditional = normalize(weighted)
            row_mass = age_sex_mass[f"{age_group} | {sex}"]
            for status in statuses:
                targets[f"{age_group} | {sex} | {status}"] = (
                    row_mass * conditional[status]
                )
    return targets


def build_country_relationship_targets(
    wpp: dict[str, Any],
    country_rates: dict[str, Any],
) -> dict[str, float]:
    """Retain the official country dimension instead of pooling it away."""

    statuses = ["single", "in relationship", "married", "divorced", "widowed"]
    total_population = sum(
        wpp["age_sex"][(country, 2025, age_group, sex)]
        for country in COUNTRY_CODES
        for age_group in AGE_GROUPS
        for sex in ("female", "male")
    )
    targets: dict[str, float] = {}
    for country in COUNTRY_CODES:
        if country not in country_rates:
            raise RuntimeError(f"UNData marital-status rows are missing for {country}.")
        for age_group in AGE_GROUPS:
            for sex in ("female", "male"):
                row_mass = (
                    wpp["age_sex"][(country, 2025, age_group, sex)] / total_population
                )
                conditional = country_rates[country][(age_group, sex)]
                for status in statuses:
                    targets[f"{country} | {age_group} | {sex} | {status}"] = (
                        row_mass * conditional.get(status, 0.0)
                    )
    return targets


def direct_ethnicity_distributions(
    ims: dict[str, Any], wpp: dict[str, Any]
) -> dict[str, dict[str, float]]:
    distributions: dict[str, dict[str, float]] = {}

    # UNData table 26, United States 2020.  The official "Asian" count is
    # split using the three Asian origin shares in UN DESA IMS 2020.
    us_asian = 19_886_049
    us_ims = ims["United States"]["origins"]
    east = us_ims.get(906, 0.0)
    south = us_ims.get(5501, 0.0)
    southeast = us_ims.get(920, 0.0)
    asian_total = east + south + southeast
    distributions["United States"] = normalize(
        {
            "White": 204_277_273,
            "Black": 41_104_200,
            "Indigenous": 3_727_135,
            "Pacific Islander": 689_966,
            "Latino": 27_915_715,
            "Mixed": 33_848_943,
            "East Asian": us_asian * east / asian_total,
            "South Asian": us_asian * south / asian_total,
            "SE Asian": us_asian * southeast / asian_total,
            "MENA": 0.0,
        }
    )

    # UNData table 26, Brazil 2022.
    distributions["Brazil"] = normalize(
        {
            "White": 88_252_121,
            "Black": 20_656_458,
            "Mixed": 92_083_286,
            "East Asian": 850_130,
            "Indigenous": 1_227_642,
        }
    )

    # UNData table 26, China 2020.  The national groups do not align with the
    # generator's regional-race taxonomy, so all nationally enumerated groups
    # stay in the broad East Asian bucket; foreigners/unknown remain Mixed.
    distributions["China"] = normalize({"East Asian": 1_408_925_641, "Mixed": 853_083})

    # UNData table 26, United Kingdom 2011.
    distributions["United Kingdom"] = normalize(
        {
            "White": 55_073_552,
            "East Asian": 433_150,
            "South Asian": 1_451_862 + 1_174_983 + 451_529,
            "Black": 1_904_684,
            "Mixed": 861_815 + 1_830_603,
        }
    )

    # Statistics Canada, 2021 Census population group and Indigenous identity.
    canada_total = 36_328_480
    visible_total = 9_639_205
    indigenous = 1_807_250
    distributions["Canada"] = normalize(
        {
            "South Asian": 2_571_400,
            "East Asian": 1_715_770 + 218_140 + 98_890,
            "Black": 1_547_870,
            "SE Asian": 957_355 + 390_340,
            "MENA": 694_015 + 360_495,
            "Latino": 580_235,
            "Mixed": 172_885 + 331_805,
            "Indigenous": indigenous,
            "White": canada_total - visible_total - indigenous,
        }
    )

    # INEGI 2020: 19.4% self-identify as Indigenous and 2.0% as
    # Afro-descendant.  These measures can overlap; normalization makes the
    # mutually-exclusive generator mapping explicit.
    distributions["Mexico"] = normalize(
        {"Indigenous": 0.194, "Black": 0.020, "Latino": 0.806}
    )

    for country, base_ethnicity in BASE_ETHNICITY.items():
        population = wpp["total"][(country, 2020)]
        origins = ims[country]["origins"]
        counts = {ethnicity: 0.0 for ethnicity in ETHNICITIES}
        counts[base_ethnicity] = max(0.0, population - origins[900])
        for origin_code, ethnicity in IMS_ORIGIN_TO_ETHNICITY.items():
            counts[ethnicity] += origins.get(origin_code, 0.0)
        distributions[country] = normalize(counts)
    return distributions


def build_ethnicity_targets(
    country_mass: dict[str, float], distributions: dict[str, dict[str, float]]
) -> dict[str, float]:
    targets: dict[str, float] = {}
    for country, mass in country_mass.items():
        for ethnicity in ETHNICITIES:
            targets[f"{country} | {ethnicity}"] = mass * distributions[country].get(
                ethnicity, 0.0
            )
    return targets
